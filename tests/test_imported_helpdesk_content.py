from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import inspect

from app.data.imported import table
from app.modules.workforce.imported_assets import insert
from app.modules.workforce.imported_helpdesk_settings import DEFAULTS, settings_table
from app.modules.workforce.imported_organization import update
from app.modules.workforce.imported_dashboards import month_offset
from app.modules.workforce.imported_assets import now
from test_imported_leave_attendance import setup_people


def test_helpdesk_drafts_settings_and_report_exports(api):
    client, app, _, headers = api
    people = setup_people(api)
    with app.state.sessions.begin() as db:
        for model in ('KnowledgeBaseFaq', 'KnowledgeBaseArticle', 'Department', 'SupportTicket'):
            table(db, model).create(db.bind)
    response = client.post('/api/ask-me/faqs', json={'question': 'Synthetic?', 'answer': 'Synthetic answer', 'isPublished': False}, headers=headers['hr'])
    assert response.status_code == 201, response.text
    assert client.get('/api/ask-me/faqs?includeDraft=1', headers=headers['employee']).json()['faqs'] == []
    assert len(client.get('/api/ask-me/faqs?includeDraft=1', headers=headers['hr']).json()['faqs']) == 1
    article = {'title': 'Synthetic article', 'excerpt': 'Synthetic excerpt', 'content': 'Synthetic content', 'tag': 'hr', 'readTimeMinutes': 2, 'isPublished': False, 'tags': 'One, one, Two'}
    response = client.post('/api/ask-me/knowledge-base', json=article, headers=headers['hr'])
    assert response.status_code == 201, response.text
    assert response.json()['article']['tags'] == ['One', 'Two']
    assert client.get('/api/ask-me/knowledge-base?includeDraft=1&includeContent=1', headers=headers['employee']).json()['articles'] == []
    path = '/api/ask-me/knowledge-base/' + response.json()['article']['id']
    assert client.patch(path, json={**article, 'isPublished': True}, headers=headers['hr']).status_code == 200
    assert 'content' not in client.get('/api/ask-me/knowledge-base?includeContent=1', headers=headers['employee']).json()['articles'][0]
    assert client.delete(path, headers=headers['employee']).status_code == 403
    response = client.get('/api/ask-me/settings', headers=headers['hr'])
    assert response.status_code == 200, response.text
    assert response.json()['settings'] == DEFAULTS
    with app.state.sessions.begin() as db:
        tbl = settings_table(db)
        assert not inspect(db.connection()).has_table(tbl.name)
        tbl.create(db.bind)
    assert client.put('/api/ask-me/settings', json=DEFAULTS, headers=headers['employee']).status_code == 403
    assert client.put('/api/ask-me/settings', json=DEFAULTS, headers=headers['hr']).status_code == 200
    with app.state.sessions.begin() as db:
        update(db, 'Employee', people['manager']['id'], {'firstName': '=2+2', 'lastName': ''})
        for number, created in enumerate((now(), datetime.combine(month_offset(now().date(), -1), datetime.min.time()))):
            insert(db, 'SupportTicket', {'ticketNumber': f'TKT-{number}', 'employeeId': people['employee']['id'], 'assignedToEmployeeId': people['manager']['id'],
                'title': 'Synthetic', 'description': 'Synthetic', 'category': 'HR', 'status': 'Open', 'priority': 'High', 'createdAt': created})
    report = client.get('/api/ask-me/reports?period=Last%20Month', headers=headers['hr'])
    assert report.status_code == 200, report.text
    assert report.json()['summary']['totalTickets'] == 1
    response = client.get('/api/ask-me/reports/export', headers=headers['hr'])
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ['Summary', 'Agent Performance', 'Category Breakdown']
    assert workbook['Agent Performance']['A2'].value == '=2+2'
    assert workbook['Agent Performance']['A2'].data_type == 's'
    response = client.get('/api/ask-me/reports/export?format=pdf', headers=headers['hr'])
    assert response.status_code == 200 and response.content.startswith(b'%PDF')
    assert client.get('/api/ask-me/reports/export', headers=headers['employee']).status_code == 403
