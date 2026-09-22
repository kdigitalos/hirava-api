"""Run with python -m app.agents.pilot --help. Output is sensitive: protect it."""
import argparse
import json
import sys
from io import BytesIO
from pathlib import Path

from pydantic import ValidationError

from app.agents.screening import create_provider, ScreeningError, ScreeningInput, build_screening_graph
from app.core.config import Settings


def extract_text(path: Path) -> str:
    if not path.is_file() or path.stat().st_size > 5 * 1024 * 1024:
        raise ScreeningError("Input must be an existing file no larger than 5 MB")
    return extract_bytes(path.read_bytes(), path.suffix.lower())


def extract_bytes(content: bytes, suffix: str) -> str:
    return extract_document(content, suffix)["text"]


def extract_document(content: bytes, suffix: str) -> dict:
    """Extract bounded S3 contents in memory; never write candidate files to disk."""
    if not content:
        raise ScreeningError("Insufficient readable text; provide a nonempty document")
    if len(content) > 5 * 1024 * 1024:
        raise ScreeningError("Input must be no larger than 5 MB")
    try:
        if suffix == ".txt":
            text = content.decode("utf-8-sig")
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted or len(reader.pages) > 30:
                raise ScreeningError("PDF must be unencrypted and at most 30 pages")
            pages = [page.extract_text() or "" for page in reader.pages]
            if any(len(page.strip()) < 80 for page in pages):
                from app.agents.ocr import local_ocr
                return local_ocr(content, suffix)
            text = "\n".join(pages)
        elif suffix in (".jpg", ".jpeg", ".png"):
            from app.agents.ocr import local_ocr
            return local_ocr(content, suffix)
        elif suffix == ".docx":
            from zipfile import ZipFile
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise ScreeningError("DOCX expanded contents exceed the extraction limit")
            from docx import Document
            doc = Document(BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs]
            paragraphs.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            text = "\n".join(paragraphs)
        else:
            raise ScreeningError("Supported files: TXT, PDF, DOCX, PNG and JPEG")
    except ScreeningError:
        raise
    except Exception:
        raise ScreeningError("Document could not be read; provide a readable TXT, PDF, DOCX, PNG or JPEG") from None
    if len(text.strip()) < 50:
        raise ScreeningError("Insufficient readable text; provide a clearer document")
    if len(text) > 60000:
        raise ScreeningError("Resume exceeds 60,000 extracted characters")
    return {"text": text, "method": "native_text", "ocr_pages": []}


def main():
    parser = argparse.ArgumentParser(description="Reviewer-operated LangGraph screening pilot")
    parser.add_argument("--resume", type=Path, required=True)
    parser.add_argument("--jd", type=Path, required=True)
    parser.add_argument("--rubric", type=Path, required=True, help="JSON containing rubric_version and criteria")
    parser.add_argument("--allow-provider-processing", action="store_true",
                        help="Confirm these documents are approved for transmission to the configured provider")
    args = parser.parse_args()
    provider = None
    try:
        if not args.allow_provider_processing:
            raise ScreeningError("Confirm approved data use with --allow-provider-processing")
        if args.rubric.stat().st_size > 100000:
            raise ScreeningError("Rubric exceeds 100 KB")
        rubric = json.loads(args.rubric.read_text(encoding="utf-8-sig"))
        source = ScreeningInput(resume_text=extract_text(args.resume),
                                job_description=extract_text(args.jd), **rubric)
        provider = create_provider(Settings())
        # Do not inherit a developer's global LangSmith tracing setting for resumes.
        from langsmith import tracing_context
        with tracing_context(enabled=False):
            result = build_screening_graph(provider).invoke({"source": source})["result"]
        print(json.dumps(result, ensure_ascii=True, indent=2))
        return 0
    except (ValidationError, TypeError):
        print("Invalid input or rubric. Check lengths, unique IDs and exact JD quotes.", file=sys.stderr)
    except (OSError, json.JSONDecodeError):
        print("Unable to read the input files or rubric JSON.", file=sys.stderr)
    except ScreeningError as exc:
        print(str(exc), file=sys.stderr)
    except ImportError:
        print('Install the AI extra: pip install -e ".[ai]"', file=sys.stderr)
    finally:
        if provider is not None:
            provider.close()
    return 1


if __name__ == "__main__":
    sys.exit(main())
