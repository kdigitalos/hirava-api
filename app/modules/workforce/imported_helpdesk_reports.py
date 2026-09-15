"""Helpdesk statistics and complete paginated PDF/Excel exports."""
from collections import Counter
from datetime import date, datetime, timedelta
from html import escape
from io import BytesIO

from fastapi import Depends, Response

from app.core.compatibility_routing import APIRouter
from app.data.database import get_db
from app.data.imported import table
from app.modules.workforce.imported_assets import admin, now, rows
from app.modules.workforce.imported_dashboards import lookup, month_offset
from app.modules.workforce.imported_employee_profile import full_name
from app.modules.workforce.imported_support import token

router = APIRouter(prefix='/api/ask-me/reports', tags=['Helpdesk reports'])
PERIODS = ('This Month', 'Last Month', 'Last 3 Months', 'This Year')


def report_data(db, period):
    period = period if period in PERIODS else 'This Month'
    today = now().date()
    start = date(today.year, 1, 1) if period == 'This Year' else month_offset(today, -1 if period == 'Last Month' else -2 if period == 'Last 3 Months' else 0)
    end = month_offset(today, 0) if period == 'Last Month' else today + timedelta(days=1)
    tbl = table(db, 'SupportTicket')
    tickets = rows(db, 'SupportTicket', tbl.c.createdAt >= datetime.combine(start, datetime.min.time()), tbl.c.createdAt < datetime.combine(end, datetime.min.time()))
    people, departments = lookup(db, 'Employee'), lookup(db, 'Department')
    durations, by_agent, categories = [], {}, Counter()
    resolved_count = 0
    for row in tickets:
        resolved = token(row['status']) in ('RESOLVED', 'CLOSED')
        resolved_count += int(resolved)
        finished = row['closedAt'] or row['resolvedAt']
        if finished:
            durations.append(max(0, (finished - row['createdAt']).total_seconds()) / 3600)
        agent = people.get(row['assignedToEmployeeId'])
        if agent:
            entry = by_agent.setdefault(agent['id'], {'name': full_name(agent), 'department': departments.get(agent['departmentId'], {}).get('name', '—'), 'assigned': 0, 'resolved': 0, 'pending': 0})
            entry['assigned'] += 1
            entry['resolved'] += int(resolved)
            entry['pending'] += int(not resolved)
        category = token(row['category'])
        categories[category if category in ('IT', 'HR') else row['category'].strip().capitalize() or 'General'] += 1
    return {'period': period, 'summary': {'totalTickets': len(tickets), 'resolvedTickets': resolved_count,
        'avgResolutionHours': sum(durations) / len(durations) if durations else 0, 'activeAgents': len(by_agent)},
        'agentPerformance': sorted(by_agent.values(), key=lambda item: (-item['assigned'], item['name'])),
        'categoryBreakdown': [{'label': label, 'tickets': count, 'total': len(tickets)} for label, count in sorted(categories.items(), key=lambda item: (-item[1], item[0]))]}


@router.get('')
def report(period: str = 'This Month', user=Depends(admin), db=Depends(get_db)):
    return report_data(db, period)


def report_sheets(report):
    summary = report['summary']
    return [('Summary', [['Metric', 'Value'], ['Period (UTC)', report['period']], ['Total Tickets', summary['totalTickets']],
        ['Resolved Tickets', summary['resolvedTickets']], ['Avg Resolution (hours)', round(summary['avgResolutionHours'], 2)], ['Active Agents', summary['activeAgents']]]),
        ('Agent Performance', [['Agent Name', 'Department', 'Assigned', 'Resolved', 'Pending', 'Resolution Rate %']] +
            [[row['name'], row['department'], row['assigned'], row['resolved'], row['pending'], round(row['resolved'] / row['assigned'] * 100) if row['assigned'] else 0] for row in report['agentPerformance']]),
        ('Category Breakdown', [['Category', 'Tickets', 'Total Tickets', 'Share %']] +
            [[row['label'], row['tickets'], row['total'], round(row['tickets'] / row['total'] * 100) if row['total'] else 0] for row in report['categoryBreakdown']])]


def excel_report(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, data in report_sheets(report):
        sheet = workbook.create_sheet(title)
        for values in data:
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                # Explicit text cells cannot execute formulas from employee names.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='4338CA')
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(60, max(16, max(len(str(row[index - 1])) for row in data) + 2))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def pdf_report(report):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=(842, 595), leftMargin=40, rightMargin=40, topMargin=35, bottomMargin=35)
    styles = getSampleStyleSheet()
    styles['BodyText'].fontSize = 9
    styles['BodyText'].leading = 12
    story = [Paragraph('HIRAVA - Helpdesk Report', styles['Title']), Paragraph(escape(report['period']) + ' | UTC', styles['Normal'])]
    for title, data in report_sheets(report):
        story.extend([Spacer(1, 12), Paragraph(title, styles['Heading2'])])
        cells = [[Paragraph(escape(str(value)), styles['BodyText']) for value in row] for row in data]
        widths = [380, 382] if len(data[0]) == 2 else [180, 180, 90, 90, 90, 132] if len(data[0]) == 6 else [270, 164, 164, 164]
        detail = Table(cells, colWidths=widths, repeatRows=1)
        detail.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eeeafa')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 7),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 7), ('LINEBELOW', (0, 0), (-1, 0), .5, colors.HexColor('#c5bddf'))]))
        story.append(detail)
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.drawString(40, 18, 'Hirava | Saved helpdesk records')
        canvas.drawRightString(802, 18, f'Page {doc.page}')
        canvas.restoreState()
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


@router.get('/export')
def export(period: str = 'This Month', format: str = 'excel', user=Depends(admin), db=Depends(get_db)):
    payload = report_data(db, period)
    pdf = format.lower() == 'pdf'
    extension = 'pdf' if pdf else 'xlsx'
    filename = 'ask-me-report-' + payload['period'].lower().replace(' ', '-') + '.' + extension
    return Response(pdf_report(payload) if pdf else excel_report(payload),
        media_type='application/pdf' if pdf else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'})
