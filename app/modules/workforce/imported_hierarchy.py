"""Manager mapping and organization hierarchy on the shared employee records."""
from io import BytesIO
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from fastapi import Depends, HTTPException, Request, Response
from pydantic import EmailStr, Field
from sqlalchemy import text

from app.core.compatibility_routing import APIRouter, CompatibilityRoute
from app.data.database import get_db
from app.data.imported import find, table
from app.modules.workforce.imported_assets import admin, insert, rows
from app.modules.workforce.imported_dashboards import avatar, lookup
from app.modules.workforce.imported_employee_profile import details, full_name
from app.modules.workforce.imported_leave import Input
from app.modules.workforce.imported_organization import update


class MappingRoute(CompatibilityRoute):
    max_body_bytes = 6 * 1024 * 1024


router = APIRouter(prefix='/api', tags=['HRMS reporting hierarchy'], route_class=MappingRoute)


def hierarchy_lock(db):
    if db.bind.dialect.name == 'postgresql':
        db.execute(text('SELECT pg_advisory_xact_lock(7310461)'))


def sorted_people(db):
    return sorted(rows(db, 'Employee'), key=lambda person: (person['lastName'] or '', person['firstName'] or ''))


def manager_label(employee):
    return full_name(employee) + ' (' + employee['employeeCode'] + ')'


@router.get('/manager-mapping')
def mappings(user=Depends(admin), db=Depends(get_db)):
    people = sorted_people(db)
    by_id = {person['id']: person for person in people}
    active = [person for person in people if person['status'] == 'ACTIVE']
    departments, titles, designations = lookup(db, 'Department'), lookup(db, 'JobTitle'), lookup(db, 'Designation')
    mapped = sum(bool(person['reportsToEmployeeId']) for person in active)
    managers = {person['reportsToEmployeeId'] for person in people}
    items = []
    for person in active:
        manager = by_id.get(person['reportsToEmployeeId'])
        items.append({'id': person['id'], 'name': full_name(person), 'email': person['email'],
            'department': departments.get(person['departmentId'], {}).get('name', '—'),
            'designation': designations.get(person['designationId'], {}).get('title') or titles.get(person['jobTitleId'], {}).get('name', '—'),
            'currentManager': manager_label(manager) if manager else None, 'currentManagerId': person['reportsToEmployeeId'], 'assignedManager': person['reportsToEmployeeId']})
    return {'stats': {'totalEmployees': len(active), 'mapped': mapped, 'unmapped': len(active) - mapped,
        'managers': sum(person['id'] in managers for person in active)}, 'rows': items,
        'managerOptions': [{'id': person['id'], 'label': manager_label(person)} for person in active]}


def assignment_error(people, employee_id, manager_id):
    person = people.get(employee_id)
    if not person or person['status'] != 'ACTIVE':
        return 'not_found'
    if manager_id:
        manager = people.get(manager_id)
        if not manager or manager['status'] != 'ACTIVE':
            return 'manager_not_found'
    if employee_id == manager_id:
        return 'self'
    seen = {employee_id}
    current = manager_id
    while current:
        if current in seen:
            return 'reporting_cycle'
        seen.add(current)
        current = people.get(current, {}).get('reportsToEmployeeId')
    return None


def set_manager(db, people, employee_id, manager_id):
    person = people[employee_id]
    metadata = {**details(person), 'manager': full_name(people[manager_id]) if manager_id else ''}
    updated = update(db, 'Employee', employee_id, {'reportsToEmployeeId': manager_id, 'employeeDetails': metadata})
    people[employee_id] = updated


def result_summary(results):
    updated = sum(result['ok'] for result in results)
    return {'results': results, 'ok': updated == len(results), 'updated': updated, 'failed': len(results) - updated}


class MappingInput(Input):
    employeeIds: list[str] = Field(min_length=1, max_length=1000)
    managerId: str | None = Field(default=None, max_length=255)


@router.post('/manager-mapping/bulk')
def bulk_mapping(body: MappingInput, user=Depends(admin), db=Depends(get_db)):
    hierarchy_lock(db)
    people = lookup(db, 'Employee')
    manager_id = body.managerId or None
    if manager_id and (manager_id not in people or people[manager_id]['status'] != 'ACTIVE'):
        raise HTTPException(422, 'Manager not found or not active')
    results = []
    for employee_id in dict.fromkeys(body.employeeIds):
        error = assignment_error(people, employee_id, manager_id)
        if not error:
            set_manager(db, people, employee_id, manager_id)
        results.append({'employeeId': employee_id, 'ok': not error, **({'error': error} if error else {})})
    return result_summary(results)


@router.get('/manager-mapping/template')
def mapping_template(user=Depends(admin)):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Manager mapping'
    sheet.append(['employee_code', 'manager_employee_code'])
    sheet.column_dimensions['A'].width = 28
    sheet.column_dimensions['B'].width = 28
    stream = BytesIO()
    workbook.save(stream)
    return Response(stream.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename="manager-mapping-template.xlsx"', 'Cache-Control': 'no-store'})


def mapping_pairs(content):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    from xml.etree.ElementTree import ParseError
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(entry.file_size for entry in entries) > 50 * 1024 * 1024:
                raise HTTPException(413, 'The expanded workbook exceeds the size limit')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.worksheets[0]
            if sheet.max_row and sheet.max_row > 1001:
                raise HTTPException(422, 'Import at most 1000 data rows at a time')
            data = sheet.iter_rows(max_row=1002, max_col=30, values_only=True)
            headers = [str(value or '').strip().lower().replace(' ', '_').replace('-', '_') for value in next(data)]
            emp = next((headers.index(key) for key in ('employee_code', 'emp_code', 'employeeid', 'code') if key in headers), None)
            mgr = next((headers.index(key) for key in ('manager_employee_code', 'manager_code', 'reports_to', 'manager') if key in headers), None)
            if emp is None or mgr is None:
                raise HTTPException(422, 'Use employee_code and manager_employee_code columns')
            pairs = []
            for row in data:
                employee = str(row[emp] or '').strip()
                manager = str(row[mgr] or '').strip()
                if not employee or employee.startswith('(') or 'example' in employee.lower():
                    continue
                if employee.startswith('=') or manager.startswith('='):
                    raise HTTPException(422, 'Employee codes must be plain text, not formulas')
                pairs.append((employee, manager))
            if len(pairs) > 1000:
                raise HTTPException(422, 'Import at most 1000 mappings at a time')
            return pairs
        finally:
            workbook.close()
    except HTTPException:
        raise
    except (BadZipFile, ValueError, KeyError, IndexError, StopIteration, InvalidFileException, ParseError):
        raise HTTPException(422, 'Upload a valid XLSX manager mapping workbook') from None


@router.post('/manager-mapping/import')
async def import_mapping(request: Request, user=Depends(admin), db=Depends(get_db)):
    async with request.form() as form:
        file = form.get('file')
        if not file or not hasattr(file, 'read'):
            raise HTTPException(422, 'Choose an XLSX file')
        content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, 'The workbook must be no larger than 5 MB')
    pairs = mapping_pairs(content)
    if not pairs:
        raise HTTPException(422, 'No data rows found')
    hierarchy_lock(db)
    people = lookup(db, 'Employee')
    codes = {person['employeeCode']: person['id'] for person in people.values() if person['status'] == 'ACTIVE'}
    results = []
    for employee_code, manager_code in pairs:
        employee_id, manager_id = codes.get(employee_code), codes.get(manager_code) if manager_code else None
        error = 'employee_not_found' if not employee_id else 'manager_not_found' if manager_code and not manager_id else assignment_error(people, employee_id, manager_id)
        if not error:
            set_manager(db, people, employee_id, manager_id)
        results.append({'employeeCode': employee_code, 'ok': not error, **({'error': error} if error else {})})
    return result_summary(results)


@router.get('/reporting-hierarchy')
def hierarchy(user=Depends(admin), db=Depends(get_db)):
    people = [person for person in sorted_people(db) if person['status'] == 'ACTIVE']
    people.sort(key=lambda person: (person['orgSortOrder'], person['firstName'], person['lastName'] or ''))
    departments, designations, titles = lookup(db, 'Department'), lookup(db, 'Designation'), lookup(db, 'JobTitle')
    nodes, parents = {}, {person['id']: person['reportsToEmployeeId'] for person in people}
    for person in people:
        photo = details(person).get('employeePhoto')
        nodes[person['id']] = {'id': person['id'], 'name': full_name(person), 'email': person['email'],
            'designation': designations.get(person['designationId'], {}).get('title') or titles.get(person['jobTitleId'], {}).get('name', '—'),
            'department': departments.get(person['departmentId'], {}).get('name', 'Unassigned'),
            'initials': avatar(person)['initials'], 'avatarColor': 'bg-blue-200 text-blue-700', 'reportsToId': person['reportsToEmployeeId'],
            'directReportCount': 0, 'children': [], **({'avatarPhoto': '/api/uploads/' + photo.lstrip('/')} if isinstance(photo, str) and photo else {})}
    roots, depth = [], 0
    for id, node in nodes.items():
        seen, current, level = {id}, parents[id], 1
        while current in nodes:
            if current in seen:
                raise HTTPException(409, 'The saved reporting hierarchy contains a cycle. Correct the manager assignments first.')
            seen.add(current)
            level += 1
            current = parents[current]
        depth = max(depth, level)
        if parents[id] in nodes:
            nodes[parents[id]]['children'].append(node)
            nodes[parents[id]]['directReportCount'] += 1
        else:
            roots.append(node)
    return {'stats': {'totalEmployees': len(people), 'departments': len(departments), 'managers': sum(bool(node['children']) for node in nodes.values()), 'hierarchyLevels': depth},
        'roots': roots, 'departments': sorted({departments[person['departmentId']]['name'] for person in people if person['departmentId'] in departments}),
        'departmentOptions': sorted([{'id': key, 'name': row['name']} for key, row in departments.items()], key=lambda row: row['name'])}


class OrderInput(Input):
    parentId: str | None = Field(default=None, max_length=255)
    orderedIds: list[str] = Field(min_length=1, max_length=1000)


@router.post('/reporting-hierarchy/reorder')
def reorder(body: OrderInput, user=Depends(admin), db=Depends(get_db)):
    hierarchy_lock(db)
    tbl = table(db, 'Employee')
    children = rows(db, 'Employee', tbl.c.status == 'ACTIVE', tbl.c.reportsToEmployeeId == (body.parentId or None))
    if len(body.orderedIds) != len(set(body.orderedIds)) or set(body.orderedIds) != {child['id'] for child in children}:
        raise HTTPException(422, 'The order must contain every sibling exactly once')
    for index, id in enumerate(body.orderedIds):
        update(db, 'Employee', id, {'orgSortOrder': index})
    return {'ok': True}


class DirectReportInput(Input):
    managerId: str = Field(min_length=1, max_length=255)
    firstName: str = Field(min_length=1, max_length=255)
    lastName: str = Field(default='', max_length=255)
    email: EmailStr | None = None
    title: str | None = Field(default=None, max_length=255)
    departmentId: str | None = Field(default=None, max_length=255)


@router.post('/reporting-hierarchy/employees', status_code=201)
def direct_report(body: DirectReportInput, user=Depends(admin), db=Depends(get_db)):
    hierarchy_lock(db)
    manager = find(db, 'Employee', body.managerId, lock=True)
    if manager['status'] != 'ACTIVE':
        raise HTTPException(422, 'The manager must be active')
    if body.departmentId:
        find(db, 'Department', body.departmentId)
    title_id = None
    if body.title:
        title = next((row for row in rows(db, 'JobTitle') if row['name'].casefold() == body.title.casefold()), None)
        title_id = (title or insert(db, 'JobTitle', {'name': body.title}))['id']
    siblings = rows(db, 'Employee', table(db, 'Employee').c.reportsToEmployeeId == body.managerId)
    record = insert(db, 'Employee', {'employeeCode': 'EMP-' + uuid4().hex[:12].upper(), 'firstName': body.firstName,
        'lastName': body.lastName, 'email': body.email, 'departmentId': body.departmentId or None, 'jobTitleId': title_id,
        'reportsToEmployeeId': body.managerId, 'status': 'ACTIVE', 'orgSortOrder': max((row['orgSortOrder'] for row in siblings), default=-1) + 1})
    return {'ok': True, 'id': record['id']}


@router.delete('/reporting-hierarchy/employees/{id}')
def remove_employee(id: str, user=Depends(admin), db=Depends(get_db)):
    find(db, 'Employee', id)
    raise HTTPException(409, 'Employee history is retained. Use the exit workflow and manager mapping instead of deleting the employee.')
